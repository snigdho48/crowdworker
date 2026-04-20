from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook


EXPECTED_HEADERS = {
    "date",
    "impressions",
    "clicks",
    "spend",
    "publisher",
}


def _norm_header(value: Any) -> str:
    return str(value or "").strip().lower()


def _parse_date(cell: Any) -> date:
    if isinstance(cell, datetime):
        return cell.date()
    if isinstance(cell, date):
        return cell
    text = str(cell).strip()
    if not text:
        raise ValueError("empty date")
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognized date: {text}")


def _parse_int(cell: Any) -> int:
    if cell is None or cell == "":
        return 0
    try:
        return int(Decimal(str(cell)))
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise ValueError(f"invalid integer: {cell}") from exc


def _parse_decimal(cell: Any) -> Decimal:
    if cell is None or cell == "":
        return Decimal("0")
    try:
        return Decimal(str(cell))
    except InvalidOperation as exc:
        raise ValueError(f"invalid number: {cell}") from exc


def parse_campaign_excel(file_obj) -> list[dict[str, Any]]:
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        raise ValueError("Excel file is empty")

    col_index = {}
    for idx, raw in enumerate(header_row):
        key = _norm_header(raw)
        if key in EXPECTED_HEADERS:
            col_index[key] = idx

    missing = EXPECTED_HEADERS - set(col_index.keys())
    if missing:
        raise ValueError(f"Missing columns: {', '.join(sorted(missing))}")

    out: list[dict[str, Any]] = []
    for row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        try:
            item = {
                "date": _parse_date(row[col_index["date"]]),
                "impressions": _parse_int(row[col_index["impressions"]]),
                "clicks": _parse_int(row[col_index["clicks"]]),
                "spend": _parse_decimal(row[col_index["spend"]]),
                "publisher": str(row[col_index["publisher"]] or "").strip(),
            }
        except (IndexError, ValueError, TypeError) as exc:
            raise ValueError(f"Row {row_num}: {exc}") from exc
        out.append(item)

    wb.close()
    return out
